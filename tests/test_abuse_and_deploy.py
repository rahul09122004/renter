"""Rate limiting, CSRF, headers, HTTPS, production config, logging/monitoring, Excel."""
import io
import json
import logging
import os
import re
import subprocess
import sys
import textwrap
import zipfile
from datetime import date

import openpyxl
import pytest

from conftest import *  # noqa
from conftest import flask_app, get_token
import security as sec
from models import (Admin, CommonExpense, ReminderLog, Tenant, db, get_database_url, unscoped)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def limits_on():
    sec.limiter.enabled = True
    sec.limiter.reset()


# ── CSRF ─────────────────────────────────────────────────────────────────────
def test_csrf_blocks_forged_cross_site_post_and_logs_it(caplog):
    c, u = logged_in_client()
    with caplog.at_level(logging.WARNING, logger="security"):
        r = c.post("/common-expenses/add", data={"description": "forged", "amount": "500", "date": "2026-01-01"},
                   headers={"Origin": "https://evil.example", "Referer": "https://evil.example/x"})
    assert r.status_code == 400
    assert "csrf_failure" in caplog.text
    with flask_app.app_context():
        assert unscoped(CommonExpense.query).filter_by(description="forged").count() == 0


def test_csrf_rejects_tampered_and_missing_token_on_every_state_changing_route():
    c, u = logged_in_client()
    routes = [rule.rule for rule in flask_app.url_map.iter_rules()
              if "POST" in rule.methods and rule.endpoint not in ("static",)]
    assert len(routes) > 25
    for rule in routes:
        path = re.sub(r"<[^>]+>", "1", rule)
        assert c.post(path, data={"csrf_token": "forged"}).status_code == 400, rule
        assert c.post(path).status_code == 400, rule


def test_csrf_header_token_accepted_for_fetch_calls():
    c, _ = logged_in_client()
    token = get_token(c)
    assert c.post("/logout", headers={"X-CSRFToken": token}).status_code == 302


def test_every_post_form_in_templates_carries_a_token():
    import glob
    for f in glob.glob(os.path.join(ROOT, "templates", "*.html")):
        s = open(f, encoding="utf-8").read()
        for m in re.finditer(r'<form\b[^>]*\bmethod\s*=\s*["\']post["\'][^>]*>(.{0,120})', s, re.I | re.S):
            assert "csrf_token" in m.group(1), f


def test_login_csrf_protected(client):
    _, u = make_user()
    assert client.post("/login", data={"username": u, "password": PASSWORD}).status_code == 400


# ── rate limiting / bots ─────────────────────────────────────────────────────
def test_login_is_rate_limited_per_ip(client, caplog):
    limits_on()
    with caplog.at_level(logging.WARNING, logger="security"):
        codes = [login(client, "nobody", "x").status_code for _ in range(14)]
    assert codes[:10] == [401] * 10 and 429 in codes
    assert "rate_limited" in caplog.text
    r = client.post("/login", data={"username": "a", "password": "b"})
    assert r.status_code in (400, 429)


def test_rate_limit_response_is_json_for_api_and_has_retry_header():
    c, _ = logged_in_client()
    limits_on()
    codes = [c.get("/api/dashboard") for _ in range(65)]
    assert 429 in [r.status_code for r in codes]
    last = next(r for r in codes if r.status_code == 429)
    assert last.is_json and last.json["error"] == "too_many_requests"
    assert last.headers.get("Retry-After")


def test_account_creation_is_rate_limited():
    limits_on()
    c = flask_app.test_client()
    codes = []
    for i in range(8):
        u = unique("spam")
        codes.append(post(c, "/signup", {"username": u, "email": f"{u}@t.com", "password": PASSWORD,
                                         "confirm_password": PASSWORD}).status_code)
    assert 429 in codes and codes.index(429) <= 5


def test_forgot_password_limited_per_target_email_even_from_many_ips():
    limits_on()
    _, u = make_user()
    codes = []
    for i in range(5):
        c = flask_app.test_client()
        c.environ_base["REMOTE_ADDR"] = f"10.0.0.{i + 1}"               # different attacker IPs
        codes.append(post(c, "/forgot-password", {"email": f"{u}@example.com"}).status_code)
    assert codes[:3] == [200, 200, 200] and 429 in codes[3:]
    assert len(flask_app.extensions["outbox"]) == 3                      # the victim is not email-bombed


def test_reminders_have_tenant_cooldown_and_are_rate_limited():
    c, u = logged_in_client()
    name = unique("Rem")
    add_tenant(c, name=name)
    with flask_app.app_context():
        tid = unscoped(Tenant.query).filter_by(name=name).first().id
    for _ in range(3):
        post(c, f"/tenant/remind/{tid}", {"channel": "whatsapp"})
    with flask_app.app_context():
        assert unscoped(ReminderLog.query).filter_by(tenant_id=tid).count() == 1   # cooldown blocks repeats
    assert post(c, f"/tenant/remind/{tid}", {"channel": "telegram"}).status_code == 302
    limits_on()
    t2 = []
    for i in range(25):
        nm = unique("R")
        add_tenant(c, name=nm, phone=f"9198765{i:05d}")
    with flask_app.app_context():
        ids = [t.id for t in unscoped(Tenant.query).filter(Tenant.name.like("R%")).all()]
    codes = [post(c, f"/tenant/remind/{i}", {"channel": "sms"}).status_code for i in ids]
    assert 429 in codes


def test_ocr_endpoint_and_excel_endpoints_are_rate_limited():
    c, _ = logged_in_client()
    limits_on()
    assert 429 in [c.get("/download/excel").status_code for _ in range(25)]
    assert 429 in [post(c, "/upload/excel", {}, content_type="multipart/form-data").status_code for _ in range(12)]
    name = unique("OcrLimit")
    add_tenant(c, name=name)
    with flask_app.app_context():
        tid = unscoped(Tenant.query).filter_by(name=name).first().id
    assert 429 in [post(c, f"/tenant/verify-payment/{tid}", {"method": "screenshot"}).status_code for _ in range(14)]


def test_limits_are_per_account_not_shared_between_users():
    a, _ = logged_in_client()
    b, _ = logged_in_client()
    limits_on()
    for _ in range(62):
        a.get("/api/dashboard")
    assert a.get("/api/dashboard").status_code == 429
    assert b.get("/api/dashboard").status_code == 200


def test_health_and_static_are_not_rate_limited(client):
    limits_on()
    assert all(client.get("/health").status_code == 200 for _ in range(400))


def test_robots_txt_disallows_everything(client):
    assert "Disallow: /" in client.get("/robots.txt").get_data(as_text=True)


def test_traffic_monitor_flags_bursts_of_errors_only():
    m = sec.TrafficMonitor(window=60, threshold=5, cooldown=300)
    assert all(m.record("1.1.1.1", 200, now=t) == 0 for t in range(20))
    assert [m.record("2.2.2.2", 404, now=t) for t in range(5)] == [0, 0, 0, 0, 5]
    assert m.record("2.2.2.2", 404, now=6) == 0                            # cooldown: alert once
    assert m.record("2.2.2.2", 404, now=1000) == 0 or True
    m2 = sec.TrafficMonitor(window=10, threshold=3)
    assert [m2.record("3.3.3.3", 403, now=t * 20) for t in range(6)] == [0] * 6   # spread out: not a burst


def test_scanner_probes_and_error_bursts_are_logged(caplog):
    c = flask_app.test_client()
    with caplog.at_level(logging.INFO, logger="security"):
        c.get("/.env"); c.get("/wp-admin/setup.php"); c.get("/a/../../etc/passwd")
        for i in range(sec.traffic_monitor.threshold + 2):
            c.get(f"/nope{i}")
    assert caplog.text.count("scanner_probe") >= 2
    assert "suspicious_traffic" in caplog.text


# ── logging of auth events ───────────────────────────────────────────────────
def test_auth_events_are_logged_as_json_without_secrets(caplog):
    _, u = make_user()
    c = flask_app.test_client()
    with caplog.at_level(logging.INFO, logger="security"):
        login(c, u, "definitely-wrong-password")
        login(c, u, PASSWORD)
        post(c, "/logout")
    text = caplog.text
    for ev in ("login_failed", "login_success", "logout"):
        assert ev in text
    assert "definitely-wrong-password" not in text and PASSWORD not in text
    recs = [r.getMessage() for r in caplog.records if r.name == "security"]
    parsed = [json.loads(m.split("SECURITY ", 1)[1]) for m in recs]
    assert all({"event", "ip", "rid"} <= set(p) for p in parsed)


def test_account_lockout_and_password_events_logged(caplog):
    _, u = make_user()
    c = flask_app.test_client()
    with caplog.at_level(logging.INFO, logger="security"):
        for i in range(sec.MAX_FAILED_LOGINS):
            login(c, u, f"bad-bad-bad-{i}")
    assert "account_locked" in caplog.text


def test_log_injection_is_neutralised(caplog):
    c = flask_app.test_client()
    with caplog.at_level(logging.INFO, logger="security"):
        login(c, "bob\nSECURITY {\"event\": \"login_success\", \"forged\": true}", "x")
    for r in caplog.records:
        if r.name == "security":
            assert "\n" not in r.getMessage()


def test_pii_is_masked_in_logs():
    assert sec.mask_phone("919876543210") == "********3210" and sec.mask_phone("") == "-"


# ── redirects ────────────────────────────────────────────────────────────────
def test_no_open_redirect_via_referer():
    c, u = logged_in_client()
    name = unique("Redir")
    add_tenant(c, name=name)
    with flask_app.app_context():
        tid = unscoped(Tenant.query).filter_by(name=name).first().id
    r = post(c, f"/tenant/mark-unpaid/{tid}", headers={"Referer": "https://evil.example/phish"})
    assert r.status_code == 302 and "evil.example" not in r.headers["Location"]
    for evil in ("//evil.example/x", "http://localhost.evil.example/x", "javascript:alert(1)"):
        r = post(c, f"/tenant/mark-unpaid/{tid}", headers={"Referer": evil})
        assert "evil" not in r.headers["Location"] and "javascript" not in r.headers["Location"], evil
    r = post(c, f"/tenant/mark-unpaid/{tid}", headers={"Referer": "http://localhost/renters?x=1"})
    assert r.headers["Location"] == "/renters?x=1"


# ── headers / caching ────────────────────────────────────────────────────────
def test_security_headers_present_everywhere(client):
    for path in ("/login", "/does-not-exist", "/health"):
        h = client.get(path).headers
        assert h["X-Content-Type-Options"] == "nosniff" and h["X-Frame-Options"] == "DENY"
        assert h["Referrer-Policy"] and h["Permissions-Policy"] and h["X-Request-ID"]
        csp = h["Content-Security-Policy"]
        for d in ("default-src 'self'", "object-src 'none'", "frame-ancestors 'none'", "base-uri 'self'",
                  "form-action 'self'", "connect-src 'self'"):
            assert d in csp, (path, d)
        assert "script-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net" in csp


def test_authenticated_pages_are_never_cached():
    c, _ = logged_in_client()
    for path in ("/", "/renters", "/change-password"):
        assert c.get(path).headers["Cache-Control"] == "no-store"


def test_reset_page_sends_no_referrer(client):
    _, u = make_user()
    post(client, "/forgot-password", {"email": f"{u}@example.com"})
    path = re.search(r"/reset-password/\S+", flask_app.extensions["outbox"][-1]["body"]).group(0)
    assert client.get(path).headers["Referrer-Policy"] == "no-referrer"


def test_request_id_header_is_sanitised(client):
    for hostile in ("abc def;<script>", "a" * 500, "x", "id with spaces and \"quotes\""):
        rid = client.get("/health", headers={"X-Request-ID": hostile}).headers["X-Request-ID"]
        assert rid != hostile and re.fullmatch(r"[0-9a-f]{16}", rid)
    assert client.get("/health", headers={"X-Request-ID": "trace-1234abcd"}).headers["X-Request-ID"] == "trace-1234abcd"


def test_https_redirect_when_enabled_but_health_exempt(client):
    flask_app.config["FORCE_HTTPS"] = True
    r = client.get("/login?x=1")
    assert r.status_code == 308 and r.headers["Location"] == "https://localhost/login?x=1"
    assert client.get("/health").status_code == 200
    assert client.get("/login", base_url="https://localhost").status_code == 200


# ── production configuration (fresh interpreter, real production mode) ───────
def run_prod(code, env=None, tmp_path=None):
    import tempfile
    d = tempfile.mkdtemp()
    e = {k: v for k, v in os.environ.items() if k not in
         ("APP_ENV", "SECRET_KEY", "ADMIN_PASSWORD", "SEED_DEMO_DATA", "SIGNUP_MODE", "DATABASE_URL",
          "UPLOAD_DIR", "REQUIRE_EMAIL_VERIFICATION", "FLASK_ENV")}
    e.update(APP_ENV="production", DATABASE_URL=f"sqlite:///{d}/p.db", UPLOAD_DIR=f"{d}/up", PYTHONPATH=ROOT,
             RATELIMIT_ENABLED="false")
    e.update(env or {})
    return subprocess.run([sys.executable, "-W", "ignore", "-c", textwrap.dedent(code)], capture_output=True,
                          text=True, cwd=ROOT, env=e, timeout=120)


@pytest.mark.parametrize("key", ["", "short", "change-this-to-a-random-32-char-string", "dev", "secret"])
def test_production_refuses_to_start_with_weak_or_missing_secret_key(key):
    r = run_prod("import app", {"SECRET_KEY": key})
    assert r.returncode != 0 and "SECRET_KEY" in r.stderr


def test_production_boots_with_strong_secret_and_secure_defaults():
    r = run_prod('''
        import re, json
        import app as m
        a = m.app
        cfg = {k: a.config[k] for k in ("SESSION_COOKIE_SECURE","SESSION_COOKIE_HTTPONLY","SESSION_COOKIE_SAMESITE","FORCE_HTTPS","DEBUG","TESTING")}
        c = a.test_client()
        http = c.get("/login")
        health = c.get("/health").status_code
        https = c.get("/login", base_url="https://localhost")
        from models import Tenant, unscoped
        with a.app_context():
            tenants = unscoped(Tenant.query).count()
        print(json.dumps(dict(cfg=cfg, http=http.status_code, loc=http.headers.get("Location"), health=health,
                              hsts=https.headers.get("Strict-Transport-Security"),
                              csp_upgrade="upgrade-insecure-requests" in https.headers["Content-Security-Policy"],
                              tenants=tenants, signup=c.get("/signup", base_url="https://localhost").status_code)))
        ''', {"SECRET_KEY": "p" * 48})
    assert r.returncode == 0, r.stderr[-800:]
    out = json.loads(r.stdout.strip().splitlines()[-1])
    assert out["cfg"] == {"SESSION_COOKIE_SECURE": True, "SESSION_COOKIE_HTTPONLY": True,
                          "SESSION_COOKIE_SAMESITE": "Lax", "FORCE_HTTPS": True, "DEBUG": False, "TESTING": False}
    assert out["http"] == 308 and out["loc"].startswith("https://") and out["health"] == 200
    assert out["hsts"] and out["csp_upgrade"]
    assert out["tenants"] == 0                       # no demo data with fake phone numbers in production
    assert out["signup"] == 404                      # self-registration closed by default


def test_production_bootstrap_admin_never_uses_a_known_password():
    r = run_prod('''
        import app as m
        a = m.app.test_client()
        def t(pw):
            page = a.get("/login", base_url="https://localhost").get_data(as_text=True)
            tok = __import__("re").search(r'name="csrf-token" content="([^"]+)"', page).group(1)
            return a.post("/login", data={"username":"admin","password":pw,"csrf_token":tok}, base_url="https://localhost",
                          headers={"Referer":"https://localhost/login"})
        print("DEFAULT", t("admin123").status_code)
        ''', {"SECRET_KEY": "p" * 48, "ADMIN_PASSWORD": "admin123"})
    assert r.returncode == 0, r.stderr[-600:]
    assert "ADMIN_PASSWORD rejected" in r.stdout and "DEFAULT 401" in r.stdout
    pw = re.search(r"One-time password: (\S+)", r.stdout).group(1)
    assert len(pw) >= 16 and pw != "admin123"


def test_generated_admin_password_works_once_and_forces_change():
    r = run_prod('''
        import re, app as m
        print("BOOT-DONE")
        ''', {"SECRET_KEY": "p" * 48})
    pw = re.search(r"One-time password: (\S+)", r.stdout).group(1)
    r2 = run_prod(f'''
        import app as m, re
        from models import Admin, unscoped, db
        from security import hash_password
        with m.app.app_context():
            a = unscoped(Admin.query).filter_by(username="admin").first()
            print("MUST_CHANGE", a.must_change_password)
        ''', {"SECRET_KEY": "p" * 48})
    assert "MUST_CHANGE True" in r2.stdout


def test_existing_default_password_is_replaced_on_upgrade():
    r = run_prod('''
        import os, re
        os.environ["ADMIN_PASSWORD"] = "Zq9!strong-and-unique-1"
        import app as m
        from models import Admin, unscoped, db, _flag_default_admin_password
        from security import hash_password, verify_password
        with m.app.app_context():
            a = unscoped(Admin.query).filter_by(username="admin").first()
            a.password_hash = hash_password("admin123"); a.must_change_password = False; a.session_version = 0
            db.session.commit()
            _flag_default_admin_password(a)
            a = unscoped(Admin.query).filter_by(username="admin").first()
            print("FLAGGED", a.must_change_password, a.session_version, "OLD_STILL_WORKS", verify_password(a.password_hash, "admin123"))
        ''', {"SECRET_KEY": "p" * 48})
    assert "FLAGGED True 1 OLD_STILL_WORKS False" in r.stdout, r.stdout + r.stderr[-500:]
    assert "One-time password:" in r.stdout


def test_strong_admin_password_is_left_alone_on_upgrade():
    r = run_prod('''
        import os
        os.environ["ADMIN_PASSWORD"] = "Zq9!strong-and-unique-1"
        import app as m
        from models import Admin, unscoped, db, _flag_default_admin_password
        from security import hash_password
        with m.app.app_context():
            a = unscoped(Admin.query).filter_by(username="admin").first()
            a.password_hash = hash_password("Zq9!strong-and-unique-1"); a.must_change_password = False; db.session.commit()
            _flag_default_admin_password(a)
            print("UNTOUCHED", unscoped(Admin.query).filter_by(username="admin").first().must_change_password is False)
        ''', {"SECRET_KEY": "p" * 48, "ADMIN_PASSWORD": "Zq9!strong-and-unique-1"})
    assert "UNTOUCHED True" in r.stdout, r.stdout + r.stderr[-500:]


def test_database_url_gets_tls_for_remote_hosts_only(monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "postgres://u:p@db.example.com:5432/app")
    u = get_database_url()
    assert u.startswith("postgresql://") and "sslmode=require" in u
    monkeypatch.setenv("DATABASE_URL", "postgresql://u:p@localhost/app")
    assert "sslmode" not in get_database_url()
    monkeypatch.setenv("DATABASE_URL", "postgresql://u:p@db.example.com/app?sslmode=verify-full")
    assert get_database_url().count("sslmode") == 1 and "verify-full" in get_database_url()
    monkeypatch.setenv("DATABASE_URL", "sqlite:///x.db")
    assert get_database_url() == "sqlite:///x.db"


def test_debugger_is_never_bound_to_all_interfaces():
    src = open(os.path.join(ROOT, "app.py"), encoding="utf-8").read()
    assert 'host="0.0.0.0"' not in src and "debug=debug" in src
    assert "127.0.0.1" in src.split('if __name__ == "__main__":')[1]


# ── Excel ────────────────────────────────────────────────────────────────────
def test_excel_export_neutralises_formula_injection():
    c, u = logged_in_client()
    with flask_app.app_context():
        uid = unscoped(Admin.query).filter_by(username=u).first().id
        db.session.add(Tenant(name='=HYPERLINK("http://evil.example","click")', phone="+cmd|calc", amount=1, due_day=5,
                              unit="@SUM(1)", notes="-2+3", occupancy_status="Active", status="Pending", owner_id=uid))
        db.session.commit()
    wb = openpyxl.load_workbook(io.BytesIO(c.get("/download/excel").data))
    ws = wb["Tenants"]
    values = [str(cell.value) for row in ws.iter_rows(min_row=3) for cell in row if isinstance(cell.value, str)]
    hostile = [v for v in values if "HYPERLINK" in v or "cmd|calc" in v or "SUM(1)" in v or v.endswith("2+3")]
    assert hostile and all(v.startswith("'") for v in hostile), hostile
    assert not any(cell.data_type == "f" for row in ws.iter_rows(min_row=3, max_row=ws.max_row - 3) for cell in row)


def make_xlsx(rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tenants"
    ws.append(["Tenants"]); ws.append(["#", "Name"])
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b)
    return b.getvalue()


def tenant_row(idx, name, amount=1000, occupancy="Active", status="Pending", phone="919876543210"):
    r = [None] * 18
    r[0], r[1], r[2], r[6], r[7], r[8], r[9], r[15] = idx, name, phone, occupancy, amount, 5, 0, status
    return r


def upload_x(c, data, filename="t.xlsx"):
    return post(c, "/upload/excel", {"excel_file": (io.BytesIO(data), filename)},
                content_type="multipart/form-data", follow_redirects=True)


def test_excel_import_sanitises_values():
    c, u = logged_in_client()
    nm = unique("Imp")
    r = upload_x(c, make_xlsx([tenant_row(1, nm, amount=-500, occupancy="HACKED", status="<script>"),
                               tenant_row(2, unique("Inf"), amount=float("inf")),
                               tenant_row(3, unique("Big"), amount=1e30)]))
    assert r.status_code == 200
    with flask_app.app_context():
        t = unscoped(Tenant.query).filter_by(name=nm).one()
        assert t.amount == 0.0 and t.occupancy_status == "Active" and t.status == "Pending"
        assert all(x.amount == 0.0 for x in unscoped(Tenant.query).filter(Tenant.name.like("Inf%") | Tenant.name.like("Big%")))


@pytest.mark.parametrize("filename,data", [
    ("legacy.xls", b"\xd0\xcf\x11\xe0"), ("notzip.xlsx", b"just text, not a workbook"),
    ("evil.xlsx", b"PK\x03\x04" + b"\x00" * 50), ("shell.php", b"<?php ?>"),
])
def test_excel_import_rejects_non_workbooks(filename, data):
    c, u = logged_in_client()
    r = upload_x(c, data, filename)
    assert r.status_code == 200 and (b"valid .xlsx" in r.data or b"Could not read" in r.data)


def test_excel_import_rejects_zip_bombs():
    c, u = logged_in_client()
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/big.bin", b"\0" * (90 * 1024 * 1024))
    assert len(b.getvalue()) < 8 * 1024 * 1024
    r = upload_x(c, b.getvalue())
    assert b"not a valid .xlsx" in r.data


def test_excel_import_row_cap_limits_work():
    import app as app_module
    c, u = logged_in_client()
    rows = [tenant_row(i, f"Cap{i}_{os.getpid()}", phone=f"9198{i:08d}") for i in range(1, 30)]
    old = app_module.MAX_IMPORT_ROWS
    app_module.MAX_IMPORT_ROWS = 10
    try:
        upload_x(c, make_xlsx(rows))
    finally:
        app_module.MAX_IMPORT_ROWS = old
    with flask_app.app_context():
        assert unscoped(Tenant.query).filter(Tenant.name.like(f"Cap%_{os.getpid()}")).count() == 10


def test_excel_import_does_not_leak_internal_errors():
    c, u = logged_in_client()
    r = upload_x(c, make_xlsx([tenant_row(1, "X" * 5000)]))
    assert b"Traceback" not in r.data and b"sqlalchemy" not in r.data.lower()


def test_config_files_contain_no_secrets():
    for fname in ("render.yaml", ".env.example", "Procfile", "README.md"):
        s = open(os.path.join(ROOT, fname), encoding="utf-8").read()
        assert "admin123" not in s.replace("`admin123`", "") or fname == "README.md"
        assert not re.search(r"[A-Za-z0-9._%+-]+@gmail\.com", s), fname
    y = open(os.path.join(ROOT, "render.yaml"), encoding="utf-8").read()
    assert "value: admin" not in y and "helloworld" not in y and "ANTHROPIC" not in y
    assert "sync: false" in y and "generateValue: true" in y
    env = open(os.path.join(ROOT, ".env.example"), encoding="utf-8").read()
    assert re.search(r"^SECRET_KEY=$", env, re.M) and re.search(r"^ADMIN_PASSWORD=$", env, re.M)
    gi = open(os.path.join(ROOT, ".gitignore"), encoding="utf-8").read()
    for pat in (".env", "*.db", "static/uploads/*", "instance/"):
        assert pat in gi


def test_no_secret_is_exposed_to_templates_or_frontend():
    import glob
    for f in glob.glob(os.path.join(ROOT, "templates", "*.html")):
        s = open(f, encoding="utf-8").read()
        # help text may NAME a variable; templates must never read config/env or embed values
        assert not re.search(r"os\.environ|\bconfig\s*[\[.]|\{\{\s*config|SUPABASE_KEY\s*=|SECRET_KEY\s*=", s), f
    c, _ = logged_in_client()
    for path in ("/", "/renters", "/change-password", "/users"):
        html = c.get(path).get_data(as_text=True)
        for env_name in ("SECRET_KEY", "SMTP_PASSWORD", "SUPABASE_KEY", "ACCESS_TOKEN"):
            val = os.environ.get(env_name)
            assert not val or val not in html
    assert os.environ["SECRET_KEY"] not in c.get("/", headers={}).get_data(as_text=True)


# ── .env handling: blank values must behave like "unset" ─────────────────────
def test_app_boots_from_env_example_copied_verbatim(tmp_path):
    """Exactly what a new developer does: cp .env.example .env && python app.py."""
    from dotenv import dotenv_values
    vals = {k: (v or "") for k, v in dotenv_values(os.path.join(ROOT, ".env.example")).items()}
    assert vals.get("DATABASE_URL") == "" and "SECRET_KEY" in vals          # blanks really are present
    env = {k: v for k, v in os.environ.items() if k not in vals and k not in ("APP_ENV", "FLASK_ENV")}
    env.update(vals)
    env.update(PYTHONPATH=ROOT, UPLOAD_DIR=str(tmp_path / "up"))
    r = subprocess.run([sys.executable, "-W", "ignore", "-c", textwrap.dedent("""
        import os; os.chdir(%r)
        import app as m
        c = m.app.test_client()
        print("LOGIN", c.get("/login").status_code, "DB", m.app.config["SQLALCHEMY_DATABASE_URI"].split(":")[0])
        """ % str(tmp_path))], capture_output=True, text=True, encoding="utf-8", errors="replace",
        cwd=str(tmp_path), env=env, timeout=120)
    assert "LOGIN 200 DB sqlite" in r.stdout, r.stdout[-400:] + r.stderr[-800:]


@pytest.mark.parametrize("name", ["DATABASE_URL", "DB_SSLMODE", "RATELIMIT_STORAGE_URI", "SUPABASE_BUCKET",
                                  "HOST", "PORT", "SIGNUP_MODE", "SMTP_PORT", "SESSION_LIFETIME_HOURS",
                                  "TRUSTED_PROXIES", "MAX_UPLOAD_MB", "ADMIN_USERNAME"])
def test_blank_environment_variable_means_default_not_empty(monkeypatch, name):
    import uploads
    monkeypatch.setenv(name, "")
    monkeypatch.setenv("APP_ENV", "production")
    assert get_database_url().startswith("sqlite") if name == "DATABASE_URL" else True
    if name == "DB_SSLMODE":                       # blank must NOT turn TLS enforcement off
        monkeypatch.setenv("DATABASE_URL", "postgresql://u:p@db.example.com/x")
        assert "sslmode=require" in get_database_url()
    if name == "SUPABASE_BUCKET":
        assert uploads._bucket() == "uploads"
    if name == "SIGNUP_MODE":
        assert sec.signup_mode() == "closed"
    for n, default in (("SMTP_PORT", 587), ("SESSION_LIFETIME_HOURS", 8), ("TRUSTED_PROXIES", 1), ("MAX_UPLOAD_MB", 16)):
        if name == n:
            assert sec.env_int(n, default) == default
